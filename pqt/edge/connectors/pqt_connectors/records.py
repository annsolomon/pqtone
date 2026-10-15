"""Reading source records and mapping them onto canonical fields (X1 B5a, ADR-037).

A source yields *raw records*: a CSV or Excel row (flat, keyed by column name), a JSON object (possibly
nested), a database row or an API object. `extract` takes only the fields a profile maps, so nothing else in
a raw record is ever looked at (ADR-033 §2). Values come back as text, except dates and times that a source
already typed (Excel cells, database columns), which stay `date` / `time` / `datetime` objects.

File formats by extension: `.csv` / `.txt`, `.xlsx`, `.json` (an array, or an object holding the records
at the profile's `path`, or under the kind's name), `.ndjson` / `.jsonl` (one record per line). JSON numbers
are read as `Decimal`, so `23.33` stays exactly 23.33.
"""
from __future__ import annotations

import csv
import io
import json
import os
import stat
import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterator

from .profile import Profile

MISSING = object()
RECORDS_SCHEMA = Path(__file__).resolve().parent / "schemas" / "records.schema.json"  # shipped with the package
MAX_FILE_BYTES = 256 * 1024 * 1024  # a day's export is far smaller; anything bigger is refused, not half-read
MAX_XLSX_UNCOMPRESSED = 64 * 1024 * 1024  # what a workbook may expand to (compressed size says little)
MAX_DIGITS = 30                            # longer JSON numbers are not amounts or quantities: refused, never expanded
OUT_OF_RANGE = "<number out of range>"


def _safe_float(literal: str) -> Decimal:
    return Decimal(literal) if len(literal) <= MAX_DIGITS + 2 and "e" not in literal.lower() else Decimal("NaN")


def _safe_int(literal: str) -> int | Decimal:
    return int(literal) if len(literal) <= MAX_DIGITS else Decimal("NaN")


def json_loads(text_: str | bytes) -> Any:
    """JSON with exact, bounded numbers: decimals as Decimal, absurd numbers as NaN (reported downstream), and
    nesting too deep for the parser as a JSONDecodeError-like ValueError."""
    try:
        return json.loads(text_, parse_float=_safe_float, parse_int=_safe_int, parse_constant=lambda c: Decimal("NaN"))
    except RecursionError:
        raise ValueError("JSON nested too deeply") from None


class SourceError(ValueError):
    """A whole source (file, sheet, document) cannot be read as the profile says."""


def resolve(record: Any, spec: str, parent: Any = None) -> Any:
    """The value `spec` names in `record`: an exact key first (CSV headers may contain dots), else a dotted
    path through objects and lists; `$.a.b` reads the outer record when rows come from `each`."""
    if spec.startswith("$."):
        return resolve(parent if parent is not None else record, spec[2:])
    if isinstance(record, dict) and spec in record:
        return record[spec]
    cur = record
    for part in spec.split("."):
        if isinstance(cur, dict) and part in cur:
            cur = cur[part]
        elif isinstance(cur, list) and part.isdigit() and int(part) < len(cur):
            cur = cur[int(part)]
        else:
            return MISSING
    return cur


def text(value: Any) -> Any:
    """Canonical text for a source value; dates and times stay typed."""
    if value is MISSING or value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, Decimal):
        if not value.is_finite() or abs(value.adjusted()) > MAX_DIGITS:
            return OUT_OF_RANGE
        return format(value, "f")
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")) or abs(value) >= 1e18:
            return OUT_OF_RANGE
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, (dict, list)):
        return json.dumps(value, sort_keys=True, default=str)  # never a valid amount or code: reported downstream
    return str(value).strip()


def extract(profile: Profile, kind: str, raw: Any, parent: Any = None) -> dict:
    cols = profile.columns[kind]
    row = {f: text(resolve(raw, src, parent)) for f, src in cols.items()}
    if kind == "dayTotals":
        row["_tenders"] = {t: text(resolve(raw, src, parent)) for t, src in profile.tenders_columns.items()}
    return row


def rows(profile: Profile, kind: str, records: Iterator[tuple[int, Any]]) -> Iterator[tuple[int, dict]]:
    """Raw records -> (record number, canonical row), expanding `each` lists."""
    each = profile.file_opts.get(kind, {}).get("each")
    for n, raw in records:
        if each:
            items = resolve(raw, each)
            if not isinstance(items, list):
                yield n, {"_error": f"record has no list at {each!r}"}
                continue
            for item in items:
                yield n, extract(profile, kind, item, raw)
        else:
            yield n, extract(profile, kind, raw)


def _json_load(text_: str) -> Any:
    return json_loads(text_)


def open_nofollow(p: Path) -> io.BufferedReader:
    """Open a source file for reading, refusing symbolic links (O_NOFOLLOW) and anything but a regular file."""
    try:
        fd = os.open(p, os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0))
    except OSError as e:
        if e.errno in (40, 62):  # ELOOP (Linux 40, macOS 62): the path is a symbolic link
            raise SourceError("is a symbolic link; only regular files are read") from None
        raise
    if not stat.S_ISREG(os.fstat(fd).st_mode):
        os.close(fd)
        raise SourceError("is not a regular file")
    return os.fdopen(fd, "rb")


def read_file(profile: Profile, kind: str, path: str | Path) -> Iterator[tuple[int, Any]]:
    """(record number, raw record) from a file; raises SourceError when the file as a whole cannot be used."""
    p = Path(path)
    ext = p.suffix.lower()
    opts = profile.file_opts.get(kind, {})
    if ext not in (".csv", ".txt", "", ".xlsx", ".json", ".ndjson", ".jsonl"):
        raise SourceError(f"unsupported file type {ext!r} (csv, txt, xlsx, json, ndjson, jsonl)")
    with open_nofollow(p) as fh:
        size = os.fstat(fh.fileno()).st_size
        if size > MAX_FILE_BYTES:
            raise SourceError(f"file is {size} bytes, more than the {MAX_FILE_BYTES} a connector reads; split the export")
        if ext in (".csv", ".txt", ""):
            yield from _csv(profile, kind, fh)
        elif ext == ".xlsx":
            yield from _xlsx(profile, kind, fh, opts.get("sheet"))
        elif ext == ".json":
            try:
                doc = _json_load(fh.read().decode("utf-8-sig"))
            except (UnicodeDecodeError, ValueError) as e:
                raise SourceError(f"not JSON: {str(e)[:120]}") from None
            records = _records_at(doc, kind, opts.get("path"))
            yield from enumerate(records, start=1)
        else:
            for i, raw_line in enumerate(fh, start=1):
                try:
                    line = raw_line.decode("utf-8-sig")
                except UnicodeDecodeError:
                    yield i, {"_error": "not UTF-8 text"}
                    continue
                if not line.strip():
                    continue
                try:
                    yield i, _json_load(line)
                except ValueError as e:
                    yield i, {"_error": f"not JSON: {str(e)[:120]}"}


def _records_at(doc: Any, kind: str, path: str | None) -> list:
    if path:
        found = resolve(doc, path)
        where = path
    elif isinstance(doc, list):
        return doc
    else:
        found = doc.get(kind, MISSING) if isinstance(doc, dict) else MISSING
        where = kind
    if not isinstance(found, list):
        raise SourceError(f"no list of records at {where!r}")
    return found


def _csv(profile: Profile, kind: str, fh) -> Iterator[tuple[int, Any]]:
    f = io.TextIOWrapper(fh, encoding=profile.csv_encoding, newline="")
    try:
        reader = csv.DictReader(f, delimiter=profile.csv_delimiter)
        try:
            _check_header(profile, kind, reader.fieldnames or [])
        except UnicodeDecodeError:
            raise SourceError(f"not {profile.csv_encoding} text") from None
        rows_ = enumerate(reader, start=2)
        while True:
            try:
                i, raw = next(rows_)
            except StopIteration:
                return
            except (UnicodeDecodeError, csv.Error) as e:
                raise SourceError(f"unreadable CSV ({type(e).__name__})") from None
            yield i, raw
    finally:
        f.detach()


def _xlsx(profile: Profile, kind: str, fh, sheet: str | None) -> Iterator[tuple[int, Any]]:
    from openpyxl import load_workbook  # MIT; only needed for Excel sources

    data = io.BytesIO(fh.read())
    try:
        with zipfile.ZipFile(data) as z:
            expanded = sum(i.file_size for i in z.infolist())
    except zipfile.BadZipFile:
        raise SourceError("not an Excel workbook (.xlsx is a zip file)") from None
    if expanded > MAX_XLSX_UNCOMPRESSED:
        raise SourceError(f"workbook expands to {expanded} bytes, more than {MAX_XLSX_UNCOMPRESSED}; save it smaller or as CSV")
    data.seek(0)
    try:
        wb = load_workbook(data, read_only=True, data_only=True)
    except Exception as e:  # KeyError, InvalidFileException, defusedxml refusals
        raise SourceError(f"not an Excel workbook ({type(e).__name__})") from None
    try:
        if sheet is not None and sheet not in wb.sheetnames:
            raise SourceError(f"no sheet {sheet!r} (sheets: {wb.sheetnames})")
        ws = wb[sheet] if sheet else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if header is None:
            raise SourceError("empty sheet")
        names = [str(h).strip() if h is not None else "" for h in header]
        _check_header(profile, kind, names)
        for i, values in enumerate(it, start=2):
            if values is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                continue
            yield i, {n: v for n, v in zip(names, values) if n}
    finally:
        wb.close()


def _check_header(profile: Profile, kind: str, header: list[str]) -> None:
    names = set(header)
    wanted = list(profile.columns[kind].values())
    if kind == "dayTotals":
        wanted += list(profile.tenders_columns.values())
    missing = [c for c in wanted if c not in names and not c.startswith("$.")]
    if missing:
        raise SourceError(f"missing column(s) {missing}")
